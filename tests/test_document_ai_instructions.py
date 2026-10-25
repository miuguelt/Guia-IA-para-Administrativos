"""Verifica la adaptación de los descargables sin alterar sus hechos ni fuentes."""
import importlib.util
import json
from pathlib import Path
import shutil
import sys
from tempfile import TemporaryDirectory
import unittest
from xml.etree import ElementTree

from docx import Document
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("document_ai_instructions", ROOT / "scripts/document_ai_instructions.py")
BUILDER = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(BUILDER)


class DocumentAIInstructionsTests(unittest.TestCase):
    def test_adaptation_preserves_sources_times_and_data_and_is_reproducible(self):
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            folder = root / "docs/Documentos"
            folder.mkdir(parents=True)
            (root / "content").mkdir()
            shutil.copy(ROOT / "content/ai-task-briefs.json", root / "content/ai-task-briefs.json")
            for name in BUILDER.FILES:
                source = ROOT / "docs/Documentos/fuentes_originales" / name
                shutil.copy(source if source.exists() else ROOT / "docs/Documentos" / name, folder / name)
            originals = {name: (folder / name).read_bytes() for name in BUILDER.FILES}
            workbook = load_workbook(folder / BUILDER.FILES[2])
            expected = {sheet.title: {cell.coordinate: (cell.value, cell.data_type, cell.number_format)
                                     for row in sheet for cell in row} for sheet in workbook}
            validations = {sheet.title: ElementTree.tostring(sheet.data_validations.to_tree()) for sheet in workbook}
            workbook.close()
            times = {}
            for name in BUILDER.FILES[:2]:
                document = Document(folder / name)
                times[name] = BUILDER.schedule_signature(document)
            reached = set()
            tracked_filename = SPEC.origin
            def track(frame, event, _argument):
                if event == "call" and frame.f_code.co_filename == tracked_filename:
                    reached.add(frame.f_code.co_name)
            sys.setprofile(track)
            try:
                self.assertEqual(BUILDER.main(root), 3)
            finally:
                sys.setprofile(None)
            for name in BUILDER.FILES:
                self.assertEqual((folder / "fuentes_originales" / name).read_bytes(), originals[name])
            for name in BUILDER.FILES[:2]:
                document = Document(folder / name)
                text = "\n".join(paragraph.text for paragraph in BUILDER.paragraphs(document))
                self.assertEqual(BUILDER.schedule_signature(document), times[name])
                for term in ["36 horas directas", "4 autónomas", "48 horas", "10 retos breves", "integrador", "resultado.json", "Antigravity", "pendiente", "material de origen"]:
                    self.assertIn(term, text)
                self.assertNotRegex(text, r"completa una ficha CRTT|respuesta manual|borrador propio|tabla a mano|piezas manualmente|Escritura individual:")
                for ident in ["reto-0", "reto-9", "reto-7", "reto-integrador"]:
                    self.assertIn(f"materiales/ia/{ident}.zip", text)
                self.assertIn("la valoración corresponde al instructor", text)
                self.assertIn("«Guarda tu evidencia»", text)
                self.assertIn("resultado.html", text)
                self.assertNotRegex(text, r"[Ee]n Documenta")
                self.assertNotIn("Evaluación retroalimentación", text)
                self.assertNotIn("propone un plan de análisis", text)
                self.assertNotIn("Luego propone borradores", text)
                self.assertNotRegex(text, r"Reto 7[ :·]+(El )?[Pp]resupuesto")
                self.assertRegex(text, r"Reto 8[ :·]+(El )?[Pp]resupuesto")
                self.assertLess(text.index("Reto 7 · Diseña el flujo"), text.index("Reto 8 · Presupuesto"))
            adapted = load_workbook(folder / BUILDER.FILES[2])
            for sheet in adapted:
                self.assertEqual(ElementTree.tostring(sheet.data_validations.to_tree()), validations[sheet.title])
                for row in sheet:
                    for cell in row:
                        if (sheet.title, cell.coordinate) not in BUILDER.WORKBOOK_NOTES:
                            self.assertEqual((cell.value, cell.data_type, cell.number_format), expected[sheet.title][cell.coordinate])
            self.assertIn("IA", adapted["Fichas de reto"]["A2"].value)
            self.assertIn("pendiente", adapted["Leeme"]["A17"].value)
            self.assertEqual(adapted["Leeme"]["A11"].value, "Reto 8")
            self.assertEqual(adapted["Leeme"]["A12"].value, "Reto integrador")
            self.assertTrue(adapted["Fichas de reto"]["A61"].value.startswith("Reto 8 · "))
            adapted.close()
            first = {name: (folder / name).read_bytes() for name in BUILDER.FILES}
            self.assertEqual(BUILDER.main(root), 3)
            self.assertEqual(first, {name: (folder / name).read_bytes() for name in BUILDER.FILES})
            self.assertTrue({"main", "paragraphs", "schedule_signature", "add_link", "adapt_docx", "adapt_workbook"}.issubset(reached), reached)
            self.assertIn("resultado.json", (root / "cartilla.txt").read_text(encoding="utf-8"))
            self.assertIn("Pida a la IA", (root / "guia_facilitador.txt").read_text(encoding="utf-8"))
            self.assertIn("La IA completa estas fichas", (root / "materiales.txt").read_text(encoding="utf-8"))

    def test_every_replacement_has_a_clear_subject_and_is_known_to_source(self):
        replacements = json.loads((ROOT / "content/document-ai-instructions.json").read_text(encoding="utf-8"))
        self.assertGreater(len(replacements), 100)
        self.assertTrue(all(before != after and after.strip() == after for before, after in replacements.items()))
        self.assertTrue(any("Pide a la IA" in value for value in replacements.values()))
        self.assertTrue(any("Pida a la IA" in value for value in replacements.values()))
        source_text = "\n".join(paragraph.text for name in BUILDER.FILES[:2]
                                for paragraph in BUILDER.paragraphs(Document(ROOT / "docs/Documentos/fuentes_originales" / name)))
        for instruction in replacements:
            self.assertIn(instruction, source_text, f"La consigna no pertenece a una fuente conservada: {instruction}")


if __name__ == "__main__":
    unittest.main()
