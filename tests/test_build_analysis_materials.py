"""Los paquetes permiten repetir el análisis sin descargar datos."""
import importlib.util
import ast
import sys
import tempfile
import csv
import hashlib
import subprocess
import re
import posixpath
from urllib.parse import unquote
from datetime import datetime
from openpyxl import load_workbook
import unittest
from pathlib import Path
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("analysis_packages", ROOT / "scripts/build-analysis-materials.py")
BUILD = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(BUILD)


class AnalysisPackagesTest(unittest.TestCase):
    def test_packages_are_complete_safe_and_reproducible(self):
        with tempfile.TemporaryDirectory() as tmp:
            target = Path(tmp)
            called = set()
            previous = sys.getprofile()
            def record(frame, event, arg):
                if event == "call" and frame.f_code.co_filename == BUILD.__file__:
                    called.add((frame.f_code.co_name, frame.f_code.co_firstlineno))
            sys.setprofile(record)
            try:
                BUILD.main(target)
            finally:
                sys.setprofile(previous)
            expected = {(n.name, n.lineno) for n in ast.walk(ast.parse(Path(BUILD.__file__).read_text(encoding="utf-8"))) if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef))}
            self.assertTrue(expected <= called, expected - called)
            print(f"Cobertura de funciones del empaquetador: {len(expected)}/{len(expected)} (100 %).")
            first = {p.name: p.read_bytes() for p in target.glob("*.zip")}
            self.assertEqual(set(first), {n + ".zip" for n in BUILD.LABS})
            for name in BUILD.LABS:
                with self.subTest(name=name), ZipFile(target / (name + ".zip")) as package:
                    paths = package.namelist()
                    self.assertEqual(len(paths), len(set(paths)))
                    self.assertFalse(any(".." in Path(p).parts or p.startswith("/") for p in paths))
                    self.assertIn("LEEME.md", paths)
                    self.assertIn("instrucciones-herramientas.md", paths)
                    self.assertIn("fichas/ficha-" + name + ".md", paths)
                    self.assertIn("guia-laboratorio.pdf", paths)
                    self.assertIn("LEEME.html", paths)
                    for task in ("lab-nasa", "lab-poblacion", "lab-clima", "lab-sismos", "lab-documentos"):
                        self.assertIn("ia/" + task + ".zip", paths)
                        for resource in ("ENCARGO-CHAT.md", "ENCARGO-IDE.md", "LEEME.html", "modelo-salida.json", "esquema-salida.json"):
                            self.assertIn("ia/" + task + "/" + resource, paths)
                        self.assertEqual(package.read("ia/" + task + ".zip"), (ROOT / "materiales/ia" / (task + ".zip")).read_bytes())
                        self.assertEqual(package.read("ia/" + task + "/esquema-salida.json"), (ROOT / "materiales/ia" / task / "esquema-salida.json").read_bytes())
                    self.assertEqual(package.read("ia/lab-nasa/fuentes/datos/exoplanetas.csv"), (ROOT / "materiales/ia/lab-nasa/fuentes/datos/exoplanetas.csv").read_bytes())
                    if name == "documentos":
                        self.assertIn("documentos/manifiesto.csv", paths)
                        originals = [p for p in paths if p.startswith('documentos/fuentes-nasa/') and p.endswith('.pdf')]
                        copies = [p for p in paths if p.startswith('ia/lab-documentos/fuentes/documentos/fuentes-nasa/') and p.endswith('.pdf')]
                        self.assertEqual(len(originals), 3)
                        self.assertEqual(len(copies), 3)
                        for original in originals:
                            self.assertEqual(package.read(original), package.read('ia/lab-documentos/fuentes/' + original))
                    else:
                        for dataset in ("exoplanetas", "poblacion", "clima", "sismos"):
                            self.assertIn("datos/" + dataset + ".csv", paths)
                        self.assertIn("codigo/analysis_data.py", paths)
                        self.assertIn("fuentes.json", paths)
                        self.assertEqual(sum(p.endswith(".ipynb") for p in paths), 1)
                        if name == "nasa":
                            extracted = target / "extraido"
                            package.extractall(extracted)
                            result = subprocess.run([sys.executable, str(extracted / "codigo/analysis_data.py"), "validar"], capture_output=True, text=True, check=True)
                            self.assertIn('4332', result.stdout)
                    self.assertNotIn("__pycache__", " ".join(paths))
                    for entry in paths:
                        if entry.endswith(('.md', '.html')):
                            body = package.read(entry).decode('utf-8')
                            links = re.findall(r'\]\(([^)]+)\)', body) if entry.endswith('.md') else re.findall(r'(?:href|src)=["\']([^"\']+)', body)
                            for href in links:
                                if not href.startswith(('https://', 'http://', '#')):
                                    linked = posixpath.normpath(posixpath.join(posixpath.dirname(entry), unquote(href.split('#')[0])))
                                    self.assertIn(linked, paths, f'{name}: {entry} -> {href}')
            BUILD.main(target)
            self.assertEqual(first, {p.name: p.read_bytes() for p in target.glob("*.zip")})

    def test_local_references_are_safe_and_missing_resources_fail(self):
        self.assertEqual(BUILD.package_path(ROOT / "materiales/ia/lab-nasa/ENCARGO-CHAT.md"), "ia/lab-nasa/ENCARGO-CHAT.md")
        self.assertEqual(BUILD.package_path(ROOT / "docs/analisis-datos/guia-laboratorio.md"), "guia-laboratorio.md")
        self.assertEqual(BUILD.package_path(ROOT / "docs/01-fundamentos.md"), "recursos/docs/01-fundamentos.md")
        with self.assertRaisesRegex(ValueError, "fuera"):
            BUILD.package_path(ROOT.parent / "no-permitido.md")
        source = ROOT / "docs/analisis-datos/guia-laboratorio.md"
        entries = {"LEEME.html": b"Indice local"}
        origins = {}
        self.assertEqual(BUILD.rewrite_reference("https://example.com/fuente", source, "guia-laboratorio.md", entries, origins), "https://example.com/fuente")
        self.assertEqual(BUILD.rewrite_reference("#seccion", source, "guia-laboratorio.md", entries, origins), "#seccion")
        self.assertEqual(BUILD.rewrite_reference("../../index.html#laboratorio-datos", source, "guia-laboratorio.md", entries, origins), "LEEME.html")
        with self.assertRaisesRegex(FileNotFoundError, "recurso"):
            BUILD.rewrite_reference("no-existe-en-el-paquete.md", source, "guia-laboratorio.md", entries, origins)

    def test_workbook_preserves_rows_values_and_dates(self):
        base = ROOT / "materiales/analisis-datos"
        workbook = load_workbook(base / "diccionario-y-datos.xlsx", read_only=True, data_only=True)
        tabs = {"Exoplanetas": "exoplanetas", "Población": "poblacion", "Clima": "clima", "Sismos": "sismos", "Santander 2024": "piramide-santander-2024"}
        self.assertEqual(set(workbook.sheetnames), {*tabs, "Diccionario", "Inicio"})
        for tab, dataset in tabs.items():
            with (base / "datos" / (dataset + ".csv")).open(encoding="utf-8-sig", newline="") as f:
                expected = list(csv.reader(f))
            actual = list(workbook[tab].iter_rows(min_row=5, max_row=len(expected) + 4, max_col=len(expected[0]), values_only=True))
            self.assertEqual(list(actual[0]), expected[0])
            self.assertEqual(len(actual), len(expected))
            for erow, arow in zip(expected[1:], actual[1:]):
                for header, source, value in zip(expected[0], erow, arow):
                    if source == "":
                        self.assertIsNone(value)
                    elif header == "fecha":
                        self.assertIsInstance(value, datetime)
                        parsed = datetime.fromisoformat(source.replace("Z", "+00:00")).replace(tzinfo=None)
                        self.assertLess(abs((parsed-value).total_seconds()), .002)
                    elif isinstance(value, (int, float)):
                        self.assertAlmostEqual(float(source), value, places=7)
                    else:
                        self.assertEqual(source, value)
        workbook.close()

    def test_document_manifest_and_included_files(self):
        base = ROOT / "materiales/analisis-datos/documentos"
        with (base / "manifiesto.csv").open(encoding="utf-8-sig", newline="") as f:
            docs = list(csv.DictReader(f))
        self.assertEqual(len(docs), 10)
        self.assertEqual(sum(int(d["paginas"]) for d in docs), 293)
        included = [d for d in docs if d["archivo_incluido"]]
        self.assertEqual(len(included), 3)
        for doc in included:
            file = base / doc["archivo_incluido"]
            if not file.is_file():
                file = ROOT / doc["archivo_incluido"]
            self.assertEqual(hashlib.sha256(file.read_bytes()).hexdigest(), doc["sha256"])


if __name__ == "__main__":
    unittest.main()
