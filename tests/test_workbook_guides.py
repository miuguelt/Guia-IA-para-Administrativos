from pathlib import Path
import unittest

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PRACTICE_BOOK = ROOT / "docs/Documentos/Materiales_practica_IA_SENA_Bogota.xlsx"
PUBLIC_DATA_BOOK = ROOT / "materiales/analisis-datos/diccionario-y-datos.xlsx"


class WorkbookGuideTests(unittest.TestCase):
    def test_practice_start_sheet_explains_excel_tools_and_safe_use(self):
        workbook = load_workbook(PRACTICE_BOOK, data_only=False, read_only=False)
        try:
            guide = workbook["Leeme"]
            text = "\n".join(
                str(cell.value)
                for row in guide.iter_rows()
                for cell in row
                if cell.value is not None
            ).casefold()

            for phrase in (
                "cómo empezar",
                "filtros",
                "tabla dinámica",
                "power query",
                "jupyter",
                "datos personales",
            ):
                with self.subTest(phrase=phrase):
                    self.assertIn(phrase, text)

            self.assertEqual(
                workbook.sheetnames,
                [
                    "Leeme",
                    "PQRS 50",
                    "PQRS complejas",
                    "Correos 50",
                    "Transcripcion",
                    "Datos SOFIA",
                    "Bienestar",
                    "Gastos Q1",
                    "Contratos",
                    "Lineamiento ficticio",
                    "Directriz",
                    "Fichas de reto",
                ],
            )
        finally:
            workbook.close()

    def test_public_data_book_has_an_excel_guide_and_formula_backed_chart(self):
        workbook = load_workbook(PUBLIC_DATA_BOOK, data_only=False, read_only=False)
        try:
            self.assertEqual(workbook.sheetnames[0], "Inicio")
            start = workbook["Inicio"]
            text = "\n".join(
                str(cell.value)
                for row in start.iter_rows()
                for cell in row
                if cell.value is not None
            ).casefold()

            for sheet_name in (
                "exoplanetas",
                "población",
                "clima",
                "sismos",
                "santander 2024",
                "power query",
                "jupyter",
            ):
                with self.subTest(sheet_or_tool=sheet_name):
                    self.assertIn(sheet_name, text)

            formulas = [
                cell.value
                for row in start.iter_rows()
                for cell in row
                if isinstance(cell.value, str) and cell.value.startswith("=")
            ]
            self.assertTrue(any("AVERAGEIFS" in formula.upper() for formula in formulas))
            self.assertTrue(any("SUMIFS" in formula.upper() for formula in formulas))
            self.assertEqual(len(start._charts), 1)
            self.assertEqual(workbook["Clima"].max_row, 371)
            self.assertEqual(workbook["Clima"]["A6"].value.year, 2024)
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
