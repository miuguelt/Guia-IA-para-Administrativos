"""Comprueba lectura, navegación y conservación de los documentos publicados."""
import importlib.util
import inspect
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from zipfile import ZipFile

from bs4 import BeautifulSoup
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("styled_documents", ROOT / "scripts/build-styled-documents.py")
BUILDER = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(BUILDER)
PDF_SPEC = importlib.util.spec_from_file_location("lab_document_pdfs", ROOT / "scripts/lab_document_pdfs.py")
PDF_BUILDER = importlib.util.module_from_spec(PDF_SPEC)
PDF_SPEC.loader.exec_module(PDF_BUILDER)


class StyledDocumentsTests(unittest.TestCase):
    covered_functions = set()
    covered_pdf_functions = set()

    def call(self, name, *args, **kwargs):
        result = getattr(BUILDER, name)(*args, **kwargs)
        self.covered_functions.add(name)
        return result

    def call_pdf(self, name, *args, **kwargs):
        result = getattr(PDF_BUILDER, name)(*args, **kwargs)
        self.covered_pdf_functions.add(name)
        return result

    def test_source_inventory_separates_current_support_and_history(self):
        sources = self.call("document_sources", ROOT)
        self.assertIn(ROOT / "docs/01-fundamentos.md", sources)
        self.assertIn(ROOT / "materiales/analisis-datos/fichas/ficha-nasa.md", sources)
        self.assertIn(ROOT / "materiales/retos/reto-00-ficha.csv", sources)
        self.assertNotIn(ROOT / "docs/architecture/README.md", sources)
        self.assertEqual(self.call("document_group", ROOT / "docs/01-fundamentos.md", "# Módulo", ROOT), "Acompañamiento del instructor")
        self.assertEqual(self.call("document_group", ROOT / "docs/guia-maestra-gemini-canva-administrativos-sena.md", "ARCHIVO HISTÓRICO", ROOT), "Antecedentes históricos")
        self.assertEqual(self.call("document_group", ROOT / "materiales/analisis-datos/fichas/ficha-nasa.md", "# NASA", ROOT), "Laboratorio de datos y documentos")
        self.assertEqual(self.call("document_group", ROOT / "docs/ejercicios/caso-01.md", "# Caso", ROOT), "Prácticas y plantillas")
        self.assertEqual(self.call("document_group", ROOT / "docs/revision-pedagogica.md", "# Revisión", ROOT), "Revisión de la guía")

    def test_reading_page_keeps_content_and_resolves_links(self):
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "docs/guia.md"
            destination = root / "docs/lectura/guia.html"
            source.parent.mkdir()
            other = root / "docs/otra.md"
            source.write_text("# Una guía\n\nTexto **importante**.\n\n## Primero\n\n[Otra](otra.md#pasos) y [web](../index.html#m01-fundamentos).\n\n![Imagen](../img/prueba.svg)\n\n| Campo | Valor |\n| --- | --- |\n| Total | 40 horas |\n\n```text\nNo inventes datos.\n```\n", encoding="utf-8")
            mapping = {source.resolve(): destination, other.resolve(): root / "docs/lectura/otra.html"}
            page = self.call("render_reading_page", source, destination, root, mapping)
            soup = BeautifulSoup(page, "html.parser")
            self.assertEqual(soup.html["lang"], "es-CO")
            self.assertEqual(soup.select_one(".document-badge").name, "span")
            self.assertEqual(soup.select_one("main h1").get_text(), "Una guía")
            self.assertEqual(soup.select_one(".document-content strong").get_text(), "importante")
            self.assertEqual(soup.select_one(".document-content table td").get_text(), "Total")
            self.assertIn('href="otra.html#pasos"', page)
            self.assertIn('href="../../index.html#m01-fundamentos"', page)
            self.assertIn('src="../../img/prueba.svg"', page)
            self.assertIn('href="#primero"', page)
            self.assertIn("No inventes datos.", soup.select_one("pre").get_text())
            self.assertIn("Ctrl+P", page)
            self.assertIn('href="../guia.md"', page)

    def test_reading_library_has_group_links_and_downloads(self):
        items = [{"title": "Guía de práctica", "group": "Guías para el aprendiz", "href": "guia.html", "source": "../guia.md"},
                 {"title": "Documento de origen", "group": "Antecedentes históricos", "href": "origen.html", "source": "../origen.md"}]
        page = self.call("render_library", items)
        soup = BeautifulSoup(page, "html.parser")
        self.assertEqual(len(soup.select(".document-card")), 2)
        self.assertIn("Antecedentes históricos", page)
        self.assertIn("Conservan su contexto de origen", page)
        self.assertIn('href="guia.html"', page)
        self.assertIn('href="../../index.html#preparacion"', page)
        self.assertEqual(len(soup.select(".download-card")), 5)

    def test_csv_sheet_lists_fields_for_the_ai_without_requesting_manual_answers(self):
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "materiales/retos/reto-00-ficha.csv"
            source.parent.mkdir(parents=True)
            source.write_text("Ficha · Reto 0\nMaterial de práctica sin registros reales.\n\nCódigo;Decisión;Justificación\n", encoding="utf-8-sig")
            page = self.call("render_csv_sheet", source, root / "docs/lectura/ficha.html", root)
            soup = BeautifulSoup(page, "html.parser")
            self.assertEqual(soup.select_one(".document-badge").name, "span")
            self.assertEqual(soup.select_one("main h1").get_text(), "Ficha · Reto 0")
            self.assertEqual([node.get_text() for node in soup.select(".worksheet-field h2")], ["Código", "Decisión", "Justificación"])
            self.assertFalse(soup.select(".answer-space,input,textarea"))
            self.assertIn('La IA genera', soup.select_one('.worksheet-guidance').get_text())
            self.assertNotIn('Escribe la respuesta', page)
            self.assertTrue(soup.select_one('a.ai-kit-reference')['href'].endswith('/materiales/ia/reto-0.zip'))
            self.assertIn("Fuente CSV", page)
            self.assertIn("Material de práctica sin registros reales.", page)
            source.write_text("Ficha\nSin campos", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "campos"):
                self.call("render_csv_sheet", source, root / "docs/lectura/ficha.html", root)

    def test_workbook_format_keeps_formulas_validations_and_data(self):
        with TemporaryDirectory() as temporary:
            path = Path(temporary) / "practica.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Material de práctica"])
            sheet.append(["No contiene registros reales."])
            sheet.append([])
            sheet.append(["Código", "Cantidad", "Total"])
            sheet.append(["A-01", 2, "=B5*3"])
            validation = DataValidation(type="whole", operator="between", formula1=1, formula2=10)
            validation.add(sheet["B5"])
            sheet.add_data_validation(validation)
            sheet.merge_cells("A1:C1")
            sheet["B5"].number_format = "0.00"
            workbook.save(path)
            before = self.call("workbook_signature", path)
            self.call("style_workbook", path)
            self.assertEqual(before, self.call("workbook_signature", path))
            styled = load_workbook(path)
            self.assertFalse(styled.active.sheet_view.showGridLines)
            self.assertEqual(styled.active.freeze_panes, "A5")
            self.assertEqual(styled.active["A4"].fill.fgColor.rgb, "001B4258")
            self.assertEqual(styled.active["B5"].number_format, "0.00")
            self.assertEqual(styled.active["C5"].value, "=B5*3")
            self.assertEqual(styled.active["B5"].border.bottom.style, "thin")
            self.assertIn("A1:C1", {str(item) for item in styled.active.merged_cells.ranges})
            self.assertIn("A2:C2", {str(item) for item in styled.active.merged_cells.ranges})
            self.assertEqual(styled.active.print_area, "'Sheet'!$A$1:$C$5")
            first = path.read_bytes()
            self.call("style_workbook", path)
            self.assertEqual(first, path.read_bytes())

    def test_docx_format_keeps_original_body_and_tables(self):
        with TemporaryDirectory() as temporary:
            path = Path(temporary) / "cartilla.docx"
            doc = Document()
            doc.add_paragraph("Guía de trabajo", "Title")
            doc.add_paragraph("Material de origen de Bogotá. Consulte la web vigente.")
            doc.add_heading("Reto 1", level=1)
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Acción"
            table.cell(0, 1).text = "Resultado"
            table.cell(1, 0).text = "Comprobar"
            table.cell(1, 1).text = "Dato pendiente"
            doc.add_page_break()
            doc.add_heading("Lista final de revisión", level=1)
            doc.add_paragraph("Comprueba la fuente.")
            doc.save(path)
            before = self.call("docx_signature", path)
            self.call("style_docx", path)
            self.assertEqual(before, self.call("docx_signature", path))
            styled = Document(path)
            self.assertEqual(styled.styles["Title"].font.color.rgb.__str__(), "111827")
            self.assertIsNotNone(styled.tables[0]._tbl.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblBorders'))
            self.assertIn("Material de origen", styled.sections[0].header.paragraphs[0].text)
            self.assertIsNotNone(styled.tables[0].rows[0]._tr.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader'))
            self.assertTrue(styled.tables[0].cell(0, 0).paragraphs[0].paragraph_format.keep_with_next)
            self.assertNotIn('w:type="page"', styled.paragraphs[-3]._p.xml)
            first = path.read_bytes()
            self.call("style_docx", path)
            self.assertEqual(first, path.read_bytes())
            self.assertEqual(before, self.call("docx_signature", path))
            self.assertEqual(len(Document(path).tables[0].cell(0, 0)._tc.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcMar/{http://schemas.openxmlformats.org/wordprocessingml/2006/main}top')), 1)

    def test_package_stabilization_is_reproducible_and_keeps_payloads(self):
        with TemporaryDirectory() as temporary:
            path = Path(temporary) / "archivo.zip"
            with ZipFile(path, "w") as archive:
                archive.writestr("document.xml", "Texto y datos intactos.")
                archive.writestr("docProps/core.xml", '<core><dcterms:modified xmlns:dcterms="http://purl.org/dc/terms/">2026-10-19T01:20:00Z</dcterms:modified></core>')
            self.call("stabilize_package", path)
            first = path.read_bytes()
            self.call("stabilize_package", path)
            self.assertEqual(first, path.read_bytes())
            with ZipFile(path) as archive:
                self.assertEqual(archive.read("document.xml").decode(), "Texto y datos intactos.")
                self.assertIn(b"2026-10-19T00:00:00Z", archive.read("docProps/core.xml"))

    def test_main_generates_complete_index_and_has_no_binary_side_effect_by_default(self):
        with TemporaryDirectory() as temporary:
            output = Path(temporary) / "lectura"
            count = self.call("main", ROOT, output)
            self.assertEqual(count, len(BUILDER.document_sources(ROOT)))
            soup = BeautifulSoup((output / "index.html").read_text(encoding="utf-8"), "html.parser")
            self.assertEqual(len(soup.select(".document-card")), count)
            for link in soup.select(".document-card .reading-link"):
                target = output / link["href"]
                self.assertTrue(target.is_file(), str(target))
            self.assertTrue((output / "manifest.json").is_file())

    def test_main_applies_opt_in_format_to_copy_only(self):
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            folder = root / "docs/Documentos"
            folder.mkdir(parents=True)
            (root / "docs/guia.md").write_text("# Guía\n\n## Paso\n\nComprueba.", encoding="utf-8")
            doc = Document()
            doc.add_paragraph("Título", "Title")
            doc.save(folder / "cartilla.docx")
            workbook = Workbook()
            workbook.active.append(["Título"])
            workbook.active.append(["Descripción"])
            workbook.active.append([])
            workbook.active.append(["Código", "Valor"])
            workbook.active.append(["A", 1])
            workbook.save(folder / "Materiales_practica_IA_SENA_Bogota.xlsx")
            self.assertEqual(self.call("main", root, root / "docs/lectura", True), 1)
            self.assertEqual(Document(folder / "cartilla.docx").styles["Title"].font.size.pt, 27)
            self.assertFalse(load_workbook(folder / "Materiales_practica_IA_SENA_Bogota.xlsx").active.sheet_view.showGridLines)

    def test_pdf_block_conversion_keeps_separation_and_order(self):
        fragment = BeautifulSoup('<li><h3>1. Formula una pregunta</h3><p>Cambia el alcance.</p><p><strong>Comprueba:</strong> la fuente.</p></li>', "html.parser")
        inline = self.call_pdf("inline", fragment.li)
        self.assertIn('pregunta</b><br/>Cambia', inline)
        self.assertIn('alcance.<br/>', inline)
        blocks = self.call_pdf("html_blocks", '<h2>Secuencia</h2><ol><li><h3>1. Pregunta</h3><p>Define el alcance.</p></li><li>Comprueba.</li></ol><table><tr><th>Dato</th></tr><tr><td>40</td></tr></table><blockquote>Conserva la fuente.</blockquote>')
        self.assertEqual([block["kind"] for block in blocks], ["h2", "li", "li", "table", "prompt"])
        self.assertIn("2. Comprueba.", blocks[2]["text"])
        self.assertEqual(blocks[3]["rows"][1], ["40"])
        with TemporaryDirectory() as temporary:
            source = Path(temporary) / "guia.md"
            source.write_text("# Guía\n\nTexto **clave**.\n\n| Campo | Valor |\n| --- | --- |\n| Duración | 40 |\n", encoding="utf-8")
            converted = self.call_pdf("md_blocks", source)
            self.assertEqual(converted[0]["text"], "Guía")
            self.assertIn("<b>clave</b>", converted[1]["text"])
            self.assertEqual(converted[2]["rows"][1][1], "40")

    def test_lab_pdfs_are_readable_complete_and_reproducible(self):
        from pypdf import PdfReader
        from reportlab.pdfgen.canvas import Canvas
        from io import BytesIO
        from types import SimpleNamespace
        styles = self.call_pdf("pdf_styles")
        self.assertEqual(styles["p"].fontSize, 10)
        blocks = self.call_pdf("document_blocks", ROOT)
        self.assertEqual(set(blocks), {"guia-laboratorio", "guia-instructor"})
        self.assertTrue(any(block["kind"] == "image" for block in blocks["guia-laboratorio"]))
        buffer = BytesIO()
        canvas = Canvas(buffer)
        self.call_pdf("page_footer", canvas, SimpleNamespace(page=3))
        canvas.showPage()
        canvas.save()
        self.assertIn("Laboratorio independiente", PdfReader(BytesIO(buffer.getvalue())).pages[0].extract_text())
        with TemporaryDirectory() as temporary:
            output = Path(temporary)
            pages = self.call_pdf("main", ROOT, output)
            self.assertGreater(pages["guia-laboratorio"], 10)
            self.assertGreaterEqual(pages["guia-instructor"], 4)
            target = output / "guia-laboratorio.pdf"
            text = " ".join(page.extract_text() for page in PdfReader(target).pages)
            self.assertIn("Cambia", text)
            self.assertNotIn("responderCambia", text)
            self.assertNotIn("copiaEntra", text)
            self.assertIn("Revisa tres criterios", PdfReader(target).pages[-1].extract_text())
            self.assertIn("resultado.json", text)
            self.assertIn("LAB-A01", text)
            self.assertIn("Cómo aplicar la rúbrica", PdfReader(output / "guia-instructor.pdf").pages[3].extract_text())
            first = target.read_bytes()
            self.call_pdf("main", ROOT, output)
            self.assertEqual(first, target.read_bytes())

    def test_zz_each_production_function_is_exercised_with_assertions(self):
        functions = {name for name, fn in inspect.getmembers(BUILDER, inspect.isfunction) if fn.__module__ == BUILDER.__name__}
        functions.update(name for name, fn in inspect.getmembers(BUILDER.DOWNLOAD_STYLES, inspect.isfunction)
                         if fn.__module__ == BUILDER.DOWNLOAD_STYLES.__name__)
        self.assertEqual(functions, self.covered_functions)
        self.assertLessEqual(len((ROOT / "scripts/build-styled-documents.py").read_text(encoding="utf-8").splitlines()), 400)
        self.assertLessEqual(len((ROOT / "scripts/document_download_styles.py").read_text(encoding="utf-8").splitlines()), 400)
        pdf_functions = {name for name, fn in inspect.getmembers(PDF_BUILDER, inspect.isfunction) if fn.__module__ == PDF_BUILDER.__name__}
        self.assertEqual(pdf_functions, self.covered_pdf_functions)


if __name__ == "__main__":
    unittest.main()
