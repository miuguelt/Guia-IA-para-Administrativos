"""
Independent Adversarial Test Oracle & Stress Harness
Identity: m_challenger_1 (critic, specialist)
Milestone: M_EXCEL_SYNC / Materials & OpenXML Verification

Tests:
1. OpenXML container, XML schemas, relationships, tableColumn-header parity, formulas, and openpyxl reader sanity.
2. 11 ZIP packages integrity, byte-by-byte parity with disk files, and challenge-materials mapping.
3. 32 CSV files standard compliance (UTF-8 BOM, CRLF, ';' delimiter, well-formedness).
4. 00-LEE-PRIMERO.txt encoding, mojibake detection, and pedagogical sections.
5. SSoT alignment between deliverables.registry.json and js/registro-entregables.js.
"""

import csv
import hashlib
import json
import os
import re
import sys
import warnings
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
EXCEL_PATH = ROOT / "docs" / "Documentos" / "Materiales_practica_IA_SENA_Bogota.xlsx"
RETOS_DIR = ROOT / "materiales" / "retos"
PAQUETES_DIR = RETOS_DIR / "paquetes"
CHALLENGE_MATERIALS_JSON = ROOT / "content" / "challenge-materials.json"
DELIVERABLES_JSON = ROOT / "deliverables.registry.json"
REGISTRO_JS = ROOT / "js" / "registro-entregables.js"
MANIFEST_JSON = ROOT / "guide.manifest.json"

NS_SPREADSHEET = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_RELS = "http://schemas.openxmlformats.org/package/2006/relationships"
NS_OFFICE_RELS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

namespaces = {
    "x": NS_SPREADSHEET,
    "r": NS_OFFICE_RELS,
    "pr": NS_RELS,
}

class TestResult:
    def __init__(self):
        self.passed = 0
        self.failed = 0
        self.errors = []
        self.observations = []

    def check(self, condition, test_name, detail=""):
        if condition:
            self.passed += 1
            print(f"  [PASS] {test_name}")
            if detail:
                self.observations.append(f"PASS: {test_name} - {detail}")
        else:
            self.failed += 1
            print(f"  [FAIL] {test_name}: {detail}")
            self.errors.append(f"{test_name}: {detail}")

def col_letter_to_index(col_str):
    exp = 0
    idx = 0
    for char in reversed(col_str.upper()):
        idx += (ord(char) - ord('A') + 1) * (26 ** exp)
        exp += 1
    return idx

def index_to_col_letter(col_idx):
    result = ""
    while col_idx > 0:
        col_idx, remainder = divmod(col_idx - 1, 26)
        result = chr(65 + remainder) + result
    return result

def parse_cell_ref(cell_ref):
    m = re.match(r"^([A-Z]+)([0-9]+)$", cell_ref)
    if not m:
        raise ValueError(f"Invalid cell ref: {cell_ref}")
    return col_letter_to_index(m.group(1)), int(m.group(2))

def parse_range_ref(range_ref):
    parts = range_ref.split(":")
    if len(parts) != 2:
        raise ValueError(f"Invalid range ref: {range_ref}")
    c1, r1 = parse_cell_ref(parts[0])
    c2, r2 = parse_cell_ref(parts[1])
    return c1, r1, c2, r2


def test_excel_openxml(res):
    print("\n" + "=" * 70)
    print("SUITE 1: ADVERSARIAL OPENXML AUDIT (docs/Documentos/Materiales_practica_IA_SENA_Bogota.xlsx)")
    print("=" * 70)

    res.check(EXCEL_PATH.is_file(), "Excel file exists on disk", str(EXCEL_PATH))
    if not EXCEL_PATH.is_file():
        return

    # 1.1 Zip container integrity
    try:
        zf = zipfile.ZipFile(EXCEL_PATH, "r")
        corrupt = zf.testzip()
        res.check(corrupt is None, "Excel ZIP container integrity", f"Corrupted file: {corrupt}")
    except Exception as e:
        res.check(False, "Excel ZIP opening", str(e))
        return

    namelist = zf.namelist()
    res.check("[Content_Types].xml" in namelist, "OpenXML [Content_Types].xml exists")
    res.check("xl/workbook.xml" in namelist, "OpenXML xl/workbook.xml exists")
    res.check("xl/_rels/workbook.xml.rels" in namelist, "OpenXML xl/_rels/workbook.xml.rels exists")

    # 1.2 Validate XML parsing of all xml and rels files
    xml_parse_errors = []
    for name in namelist:
        if name.endswith(".xml") or name.endswith(".rels"):
            try:
                data = zf.read(name)
                ET.fromstring(data)
            except Exception as e:
                xml_parse_errors.append((name, str(e)))
    res.check(len(xml_parse_errors) == 0, "All XML and rels files are well-formed", f"Errors: {xml_parse_errors}")

    # 1.3 Load shared strings
    shared_strings = []
    if "xl/sharedStrings.xml" in namelist:
        try:
            ss_tree = ET.fromstring(zf.read("xl/sharedStrings.xml"))
            for si in ss_tree.findall(f"{{{NS_SPREADSHEET}}}si"):
                t_el = si.find(f"{{{NS_SPREADSHEET}}}t")
                if t_el is not None and t_el.text:
                    shared_strings.append(t_el.text)
                else:
                    # Could be rich text (r/t)
                    text_parts = [t.text for t in si.findall(f".//{{{NS_SPREADSHEET}}}t") if t.text]
                    shared_strings.append("".join(text_parts))
        except Exception as e:
            res.check(False, "Parse sharedStrings.xml", str(e))

    # 1.4 Workbook and sheet relations
    wb_tree = ET.fromstring(zf.read("xl/workbook.xml"))
    wb_rels_tree = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))

    sheet_rids = {}
    for sheet_el in wb_tree.findall(f".//{{{NS_SPREADSHEET}}}sheet"):
        s_name = sheet_el.get("name")
        s_rid = sheet_el.get(f"{{{NS_OFFICE_RELS}}}id")
        sheet_rids[s_rid] = s_name

    rel_targets = {}
    for rel_el in wb_rels_tree.findall(f".//{{{NS_RELS}}}Relationship"):
        r_id = rel_el.get("Id")
        target = rel_el.get("Target")
        rel_targets[r_id] = target

    sheet_to_xml = {}
    for r_id, s_name in sheet_rids.items():
        target = rel_targets.get(r_id, "")
        norm_target = target.lstrip("/")
        if not norm_target.startswith("xl/"):
            norm_target = "xl/" + norm_target
        sheet_to_xml[s_name] = norm_target

    res.check(len(sheet_to_xml) == 12, "Workbook declares exactly 12 sheets", f"Found: {list(sheet_to_xml.keys())}")

    # 1.5 Cross-reference tables and worksheets
    # Find all tables in zip
    table_files = [n for n in namelist if n.startswith("xl/tables/table") and n.endswith(".xml")]
    res.check(len(table_files) == 10, "Workbook contains exactly 10 table files", f"Count: {len(table_files)}")

    # Map each sheet to its tables via sheet rels
    table_to_sheet = {}
    for s_name, s_xml in sheet_to_xml.items():
        s_dir = os.path.dirname(s_xml)
        s_file = os.path.basename(s_xml)
        rels_path = f"{s_dir}/_rels/{s_file}.rels"
        if rels_path in namelist:
            s_rels_tree = ET.fromstring(zf.read(rels_path))
            for rel in s_rels_tree.findall(f".//{{{NS_RELS}}}Relationship"):
                r_type = rel.get("Type", "")
                if r_type.endswith("table"):
                    target = rel.get("Target", "")
                    norm_target = target.lstrip("/")
                    if not norm_target.startswith("xl/"):
                        if target.startswith("../"):
                            norm_target = "xl/" + target.split("../")[-1]
                        else:
                            norm_target = f"{s_dir}/{target}"
                    full_target = os.path.normpath(norm_target).replace("\\", "/")
                    table_to_sheet[full_target] = (s_name, s_xml)

    res.check(len(table_to_sheet) == 10, "All 10 tables are mapped to parent worksheets via sheet rels")

    # 1.6 Verify Column Name Parity for EVERY table against its worksheet header row!
    total_columns_checked = 0
    discrepancies = []
    duplicate_columns = []

    for t_file in sorted(table_files):
        t_tree = ET.fromstring(zf.read(t_file))
        t_name = t_tree.get("name")
        t_ref = t_tree.get("ref")
        sheet_info = table_to_sheet.get(t_file)
        if not sheet_info:
            discrepancies.append((t_file, "Unmapped to worksheet"))
            continue

        s_name, s_xml = sheet_info
        s_tree = ET.fromstring(zf.read(s_xml))

        # Parse table columns in table XML
        table_cols = []
        col_names_seen = set()
        for col_el in t_tree.findall(f".//{{{NS_SPREADSHEET}}}tableColumn"):
            c_name = col_el.get("name")
            c_id = col_el.get("id")
            table_cols.append((c_id, c_name))
            if c_name in col_names_seen:
                duplicate_columns.append((t_file, c_name))
            col_names_seen.add(c_name)

        c1, r1, c2, r2 = parse_range_ref(t_ref)
        expected_col_count = c2 - c1 + 1
        if len(table_cols) != expected_col_count:
            discrepancies.append((t_file, f"Column count mismatch: {len(table_cols)} cols in XML vs {expected_col_count} in ref {t_ref}"))

        # Now extract the cells in worksheet row r1 from col c1 to col c2
        header_row_el = None
        for row_el in s_tree.findall(f".//{{{NS_SPREADSHEET}}}row"):
            if int(row_el.get("r", 0)) == r1:
                header_row_el = row_el
                break

        if header_row_el is None:
            discrepancies.append((t_file, f"Header row {r1} not found in sheet {s_name} ({s_xml})"))
            continue

        sheet_cells = {}
        for c_el in header_row_el.findall(f"{{{NS_SPREADSHEET}}}c"):
            c_ref = c_el.get("r")
            c_type = c_el.get("t")
            val_el = c_el.find(f"{{{NS_SPREADSHEET}}}v")
            val = val_el.text if val_el is not None else ""
            if c_type == "s" and val.isdigit():
                idx = int(val)
                text_val = shared_strings[idx] if idx < len(shared_strings) else f"[OUT_OF_BOUNDS_{idx}]"
            elif c_type == "inlineStr":
                is_t = c_el.find(f".//{{{NS_SPREADSHEET}}}t")
                text_val = is_t.text if is_t is not None else ""
            else:
                text_val = val
            sheet_cells[c_ref] = text_val

        # Compare table column names with sheet cells
        for idx, (c_id, col_name) in enumerate(table_cols):
            total_columns_checked += 1
            col_letter = index_to_col_letter(c1 + idx)
            cell_ref = f"{col_letter}{r1}"
            cell_text = sheet_cells.get(cell_ref, "")

            if col_name != cell_text:
                discrepancies.append((
                    t_file,
                    f"Sheet '{s_name}' cell {cell_ref}: cell text '{cell_text}' != table column name '{col_name}'"
                ))

    res.check(len(duplicate_columns) == 0, "No duplicate column names inside any table", str(duplicate_columns))
    res.check(
        len(discrepancies) == 0 and total_columns_checked >= 40,
        f"All {total_columns_checked} table column names strictly match sheet header cells",
        f"Discrepancies: {discrepancies}"
    )

    # 1.7 Formula verification in sheet7 (Bienestar)
    bienestar_xml = sheet_to_xml.get("Bienestar")
    res.check(bienestar_xml is not None, "Sheet 'Bienestar' found in workbook")
    if bienestar_xml:
        b_tree = ET.fromstring(zf.read(bienestar_xml))
        formulas_found = {}
        for c_el in b_tree.findall(f".//{{{NS_SPREADSHEET}}}c"):
            f_el = c_el.find(f"{{{NS_SPREADSHEET}}}f")
            if f_el is not None and f_el.text:
                formulas_found[c_el.get("r")] = f_el.text

        # Verify E5:E14 formulas (Tasa de deserción = C/B)
        bienestar_formula_ok = True
        for row in range(5, 15):
            ref = f"E{row}"
            expected_f = f"C{row}/B{row}"
            actual_f = formulas_found.get(ref, "")
            if actual_f != expected_f and not (row > 5 and ref not in formulas_found and "E5" in formulas_found):
                bienestar_formula_ok = False
                res.observations.append(f"Bienestar formula at {ref}: {actual_f} (expected {expected_f})")
        res.check(bienestar_formula_ok, "Bienestar E5:E14 formulas integrity (relative references =C{n}/B{n})")

    # 1.8 Check for calculation/formula errors stored in cells (#REF!, #VALUE!, #NAME?, #N/A)
    error_cells = []
    for s_name, s_xml in sheet_to_xml.items():
        s_tree = ET.fromstring(zf.read(s_xml))
        for c_el in s_tree.findall(f".//{{{NS_SPREADSHEET}}}c"):
            v_el = c_el.find(f"{{{NS_SPREADSHEET}}}v")
            if v_el is not None and v_el.text:
                if v_el.text.startswith("#") and v_el.text in {"#REF!", "#VALUE!", "#NAME?", "#N/A", "#NULL!", "#NUM!", "#DIV/0!"}:
                    error_cells.append((s_name, c_el.get("r"), v_el.text))
    res.check(len(error_cells) == 0, "No stored formula error values in any worksheet", str(error_cells))

    # 1.9 Test openpyxl loading without warnings or errors
    try:
        from openpyxl import load_workbook
        with warnings.catch_warnings(record=True) as caught_warnings:
            warnings.simplefilter("always")
            wb = load_workbook(EXCEL_PATH, data_only=False)
            warn_msgs = [str(w.message) for w in caught_warnings]
            critical_warns = [w for w in warn_msgs if "table" in w.lower() or "corrupt" in w.lower() or "repair" in w.lower()]
            res.check(len(critical_warns) == 0, "openpyxl loads workbook with 0 table/corruption warnings", str(warn_msgs))
            
            # Check sheet count and names
            res.check(len(wb.sheetnames) == 12, "openpyxl reads 12 sheets", str(wb.sheetnames))
            
            # Count tables via openpyxl
            total_wb_tables = sum(len(ws.tables) for ws in wb.worksheets)
            res.check(total_wb_tables == 10, "openpyxl detects all 10 tables across worksheets")
            wb.close()
    except Exception as e:
        res.check(False, "openpyxl load_workbook execution", str(e))


def test_zip_packages_and_parity(res):
    print("\n" + "=" * 70)
    print("SUITE 2: ADVERSARIAL ZIP PACKAGES & BYTE PARITY AUDIT (materiales/retos/paquetes/)")
    print("=" * 70)

    res.check(PAQUETES_DIR.is_dir(), "Directory materiales/retos/paquetes exists")
    if not PAQUETES_DIR.is_dir():
        return

    expected_packages = [
        "reto-0.zip", "reto-1.zip", "reto-2.zip", "reto-3.zip", "reto-4.zip",
        "reto-5.zip", "reto-6.zip", "reto-7.zip", "reto-8.zip", "reto-9.zip",
        "reto-integrador.zip"
    ]
    actual_packages = sorted(p.name for p in PAQUETES_DIR.glob("*.zip"))
    res.check(actual_packages == sorted(expected_packages), "Exactly 11 challenge packages exist", f"Found: {actual_packages}")

    # Load challenge-materials.json to verify package completeness
    res.check(CHALLENGE_MATERIALS_JSON.is_file(), "content/challenge-materials.json exists")
    challenge_materials = {}
    if CHALLENGE_MATERIALS_JSON.is_file():
        challenge_materials = json.loads(CHALLENGE_MATERIALS_JSON.read_text(encoding="utf-8"))

    total_zip_files_audited = 0
    byte_parity_mismatches = []
    missing_files = []
    corrupt_zips = []
    readme_issues = []

    mojibake_pattern = re.compile(r"Ã[¡¢£¤¥¦§¨©ª«¬­®¯°±²³´µ¶·¸¹º»¼½¾¿ÁÂÃÄÅÆÇÈÉÊËÌÍÎÏÐÑÒÓÔÕÖ×ØÙÚÛÜÝÞßàáâãäåæçèéêëìíîïðñòóôõö÷øùúûüýþÿ]|â€|Â|\ufffd")

    for pkg_name in actual_packages:
        pkg_path = PAQUETES_DIR / pkg_name
        challenge_id = pkg_name.replace(".zip", "")
        try:
            zf = zipfile.ZipFile(pkg_path, "r")
            test_res = zf.testzip()
            if test_res is not None:
                corrupt_zips.append((pkg_name, test_res))
                continue
            
            namelist = zf.namelist()

            # 2.1 Audit 00-LEE-PRIMERO.txt
            if "00-LEE-PRIMERO.txt" not in namelist:
                readme_issues.append((pkg_name, "Missing 00-LEE-PRIMERO.txt"))
            else:
                raw_readme = zf.read("00-LEE-PRIMERO.txt")
                try:
                    text_readme = raw_readme.decode("utf-8")
                    if len(text_readme) < 150:
                        readme_issues.append((pkg_name, f"00-LEE-PRIMERO.txt suspiciously short ({len(text_readme)} chars)"))
                    
                    # Mojibake check
                    mojibake_matches = mojibake_pattern.findall(text_readme)
                    if mojibake_matches:
                        readme_issues.append((pkg_name, f"Mojibake detected in 00-LEE-PRIMERO.txt: {set(mojibake_matches)}"))

                    # Required sections
                    for sec in ["RETO ·", "CÓMO TRABAJAR", "ARCHIVOS INCLUIDOS", "CONSIGNA Y PASOS DEL RETO"]:
                        if sec not in text_readme:
                            readme_issues.append((pkg_name, f"Missing section '{sec}' in 00-LEE-PRIMERO.txt"))

                    if challenge_id in {"reto-1", "reto-2", "reto-3"}:
                        if "CRTT" not in text_readme:
                            readme_issues.append((pkg_name, f"Missing CRTT instructions in {challenge_id} 00-LEE-PRIMERO.txt"))

                except UnicodeDecodeError as e:
                    readme_issues.append((pkg_name, f"00-LEE-PRIMERO.txt is not valid UTF-8: {e}"))

            # 2.2 Byte-by-byte parity check of included materials
            expected_specs = challenge_materials.get(challenge_id, [])
            expected_filenames = {spec["href"].rsplit("/", 1)[-1]: spec for spec in expected_specs}

            for member in namelist:
                if member == "00-LEE-PRIMERO.txt":
                    continue
                total_zip_files_audited += 1

                # Member must match an expected file for this challenge
                if member not in expected_filenames:
                    missing_files.append((pkg_name, f"Unexpected member '{member}' in zip"))

                # Resolve disk file
                disk_file = RETOS_DIR / member
                if not disk_file.is_file():
                    missing_files.append((pkg_name, f"File '{member}' from zip not found on disk at {disk_file}"))
                    continue

                # Byte comparison & SHA-256
                zip_bytes = zf.read(member)
                disk_bytes = disk_file.read_bytes()

                if len(zip_bytes) != len(disk_bytes):
                    byte_parity_mismatches.append((
                        pkg_name, member,
                        f"Size mismatch: zip={len(zip_bytes)} bytes vs disk={len(disk_bytes)} bytes"
                    ))
                elif zip_bytes != disk_bytes:
                    zip_sha = hashlib.sha256(zip_bytes).hexdigest()
                    disk_sha = hashlib.sha256(disk_bytes).hexdigest()
                    byte_parity_mismatches.append((
                        pkg_name, member,
                        f"Hash mismatch: zip={zip_sha[:12]} vs disk={disk_sha[:12]}"
                    ))

            # Verify that every expected file is present in zip
            for exp_fname in expected_filenames:
                if exp_fname not in namelist:
                    missing_files.append((pkg_name, f"Expected file '{exp_fname}' missing from zip"))

            zf.close()
        except Exception as e:
            corrupt_zips.append((pkg_name, str(e)))

    res.check(len(corrupt_zips) == 0, "All 11 ZIP archives pass integrity testzip()", str(corrupt_zips))
    res.check(len(readme_issues) == 0, "All 00-LEE-PRIMERO.txt are UTF-8, clean, without mojibake and with all sections", str(readme_issues))
    res.check(
        len(byte_parity_mismatches) == 0 and total_zip_files_audited >= 25,
        f"Byte-by-byte exact parity between {total_zip_files_audited} files inside ZIPs and disk files",
        str(byte_parity_mismatches)
    )
    res.check(len(missing_files) == 0, "Full bidirectional match between ZIP contents and challenge-materials.json", str(missing_files))


def test_csv_standards(res):
    print("\n" + "=" * 70)
    print("SUITE 3: ADVERSARIAL CSV FORMAT STANDARDS (materiales/retos/*.csv)")
    print("=" * 70)

    csv_files = sorted(RETOS_DIR.glob("*.csv"))
    res.check(len(csv_files) == 32, f"Exactly 32 CSV files exist in materiales/retos/ (found {len(csv_files)})")

    bom_failures = []
    crlf_failures = []
    delimiter_failures = []
    csv_parse_errors = []

    BOM_UTF8 = b"\xef\xbb\xbf"

    for csv_path in csv_files:
        name = csv_path.name
        raw_bytes = csv_path.read_bytes()

        # 3.1 UTF-8 BOM check
        if not raw_bytes.startswith(BOM_UTF8):
            bom_failures.append(name)

        # 3.2 CRLF newline check: must have \r\n and NO standalone \n (without preceding \r)
        # Check if \r\n exists
        if b"\r\n" not in raw_bytes and len(raw_bytes) > 0:
            crlf_failures.append((name, "No CRLF found"))
        else:
            # Check for isolated LF
            # Remove all \r\n, if any \n remains, it's a standalone LF
            stripped = raw_bytes.replace(b"\r\n", b"")
            if b"\n" in stripped:
                crlf_failures.append((name, "Standalone LF detected"))
            if b"\r" in stripped:
                crlf_failures.append((name, "Standalone CR detected"))

        # 3.3 Delimiter & parsing check
        try:
            text = raw_bytes.decode("utf-8-sig")
            lines = text.split("\r\n")
            reader = list(csv.reader(lines, delimiter=";"))
            
            # Check if semicolon is actually present in data rows
            # Some files have title rows with 1 column, but subsequent table rows must have multiple columns separated by ;
            max_cols = max(len(row) for row in reader if row) if reader else 0
            if max_cols < 2 and len(reader) > 2:
                delimiter_failures.append((name, f"Max columns detected is {max_cols} (expected multiple columns separated by ;)"))
        except Exception as e:
            csv_parse_errors.append((name, str(e)))

    res.check(len(bom_failures) == 0, "All 32 CSV files start with UTF-8 BOM (\\xef\\xbb\\xbf)", str(bom_failures))
    res.check(len(crlf_failures) == 0, "All 32 CSV files strictly use Windows CRLF (\\r\\n) line endings", str(crlf_failures))
    res.check(len(delimiter_failures) == 0, "All 32 CSV files correctly format tabular data with ';' delimiter", str(delimiter_failures))
    res.check(len(csv_parse_errors) == 0, "All 32 CSV files parse without csv syntax errors", str(csv_parse_errors))


def test_ssot_and_editorial(res):
    print("\n" + "=" * 70)
    print("SUITE 4: SSOT DELIVERABLES REGISTRY & EDITORIAL ALIGNMENT")
    print("=" * 70)

    # 4.1 Check existence of deliverables.registry.json and js/registro-entregables.js
    res.check(DELIVERABLES_JSON.is_file(), "deliverables.registry.json exists")
    res.check(REGISTRO_JS.is_file(), "js/registro-entregables.js exists")
    if not (DELIVERABLES_JSON.is_file() and REGISTRO_JS.is_file()):
        return

    # 4.2 Compare JSON in registry with JS export
    registry_data = json.loads(DELIVERABLES_JSON.read_text(encoding="utf-8"))
    js_text = REGISTRO_JS.read_text(encoding="utf-8")

    # js_text should be: export const REGISTRO_ENTREGABLES = ... or export default ...
    # Extract JSON string from JS file
    m = re.search(r"=\s*(\{.*\});?\s*$", js_text, re.DOTALL)
    if m:
        try:
            js_data = json.loads(m.group(1))
            res.check(registry_data == js_data, "deliverables.registry.json and js/registro-entregables.js are 100% data-identical")
        except Exception as e:
            res.check(False, "Parse JS exported object as JSON", str(e))
    else:
        res.check(False, "Extract JSON structure from js/registro-entregables.js")

    # 4.3 Check Rule 5.0.1 compliance: no redundant adjective 'ficticio' / 'sintético' in deliverables
    deliv_text = DELIVERABLES_JSON.read_text(encoding="utf-8")
    redundant_ficticio = re.findall(r"\b(caso|ejemplo|dato[s]?|persona|archivo|registro[s]?)\s+ficticio[s]?\b", deliv_text, re.IGNORECASE)
    all_sintetico = re.findall(r"sint[eé]tico", deliv_text, re.IGNORECASE)
    
    # Verify that all occurrences of 'ficticio' are strictly references to the proper noun 'Lineamiento ficticio' or 'lineamiento-ficticio.csv'
    all_ficticio = re.findall(r"ficticio", deliv_text, re.IGNORECASE)
    allowed_ficticio = re.findall(r"lineamiento[- ]ficticio", deliv_text, re.IGNORECASE)

    res.check(
        len(redundant_ficticio) == 0 and len(all_sintetico) == 0 and len(all_ficticio) == len(allowed_ficticio),
        f"deliverables.registry.json strictly complies with Rule 5.0.1 (0 redundant adjectives, all {len(all_ficticio)} 'ficticio' are proper nouns of 'Lineamiento ficticio')",
        f"Redundant adjectives: {redundant_ficticio}, Sintetico: {all_sintetico}, Unaccounted ficticio: {len(all_ficticio) - len(allowed_ficticio)}"
    )


def main():
    print("*" * 70)
    print("EMPIRICAL ADVERSARIAL ORACLE HARNESS — m_challenger_1")
    print("*" * 70)

    res = TestResult()
    test_excel_openxml(res)
    test_zip_packages_and_parity(res)
    test_csv_standards(res)
    test_ssot_and_editorial(res)

    print("\n" + "=" * 70)
    print(f"VERIFICATION SUMMARY: {res.passed} PASSED, {res.failed} FAILED")
    print("=" * 70)

    if res.failed > 0:
        print("\nFAILURES ENCOUNTERED:")
        for err in res.errors:
            print(f"  - {err}")
        print("\nVERDICT: REQUEST_CHANGES")
        sys.exit(1)
    else:
        print("\nALL ADVERSARIAL ORACLES PASSED EMPIRICALLY.")
        print("VERDICT: APPROVE")
        sys.exit(0)

if __name__ == "__main__":
    main()
