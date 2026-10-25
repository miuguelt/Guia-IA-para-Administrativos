"""Descarga explícita, validación offline y análisis de cuatro fuentes oficiales.

Uso desde la raíz del repositorio:
  python materiales/analisis-datos/codigo/analysis_data.py validar
  python materiales/analisis-datos/codigo/analysis_data.py analizar
  python materiales/analisis-datos/codigo/analysis_data.py figuras
  python materiales/analisis-datos/codigo/analysis_data.py refrescar

Las acciones «refrescar» y «refrescar-dane» acceden a la red. Las figuras requieren
Matplotlib y «refrescar-dane» requiere openpyxl; las demás operaciones usan la
biblioteca estándar. Los originales se conservan.
"""
import argparse
import csv
from datetime import datetime, timezone
import hashlib
import io
import json
import math
from pathlib import Path
import re
import statistics
from urllib.parse import urlencode
from urllib.request import Request, urlopen

BASE = Path(__file__).resolve().parents[1]
DANE_URL = 'https://www.dane.gov.co/files/censo2018/proyecciones-de-poblacion/Departamental/PPED-AreaSexoEdadDep-2018-2050_VP.xlsx'
PLANET_QUERY = ("select pl_name,hostname,discoverymethod,pl_rade,pl_orbper,st_teff,st_rad "
                "from ps where default_flag=1 and discoverymethod='Transit' "
                "and disc_year<=2024 and pl_rade>0 and pl_orbper>0 "
                "and pl_radelim=0 and pl_orbperlim=0 order by pl_name")
DATASETS = {
    'exoplanetas': {
        'url': 'https://exoplanetarchive.ipac.caltech.edu/TAP/sync?' + urlencode({'query': PLANET_QUERY, 'format': 'json'}),
        'documentation': 'https://exoplanetarchive.ipac.caltech.edu/docs/API_TD_columns.html',
        'fields': ['pl_name', 'hostname', 'discoverymethod', 'pl_rade', 'pl_orbper', 'st_teff', 'st_rad'],
        'numeric': ['pl_rade', 'pl_orbper', 'st_teff', 'st_rad'], 'nullable': ['st_teff', 'st_rad'],
        'key': ['pl_name'],
        'units': {'pl_name': 'identificador', 'hostname': 'nombre de anfitriona', 'discoverymethod': 'categoría',
                  'pl_rade': 'radios terrestres', 'pl_orbper': 'días', 'st_teff': 'K', 'st_rad': 'radios solares'},
        'license': 'Datos de acceso público. Citar NASA Exoplanet Archive y las referencias originales; la página de atribución no declara una licencia CC única.',
        'license_url': 'https://exoplanetarchive.ipac.caltech.edu/docs/acknowledge.html',
        'filters': PLANET_QUERY,
        'transformation': 'Se conservan las siete columnas consultadas. Se ordenan por nombre. Nulos estelares se escriben como celdas vacías; no se imputan.',
        'limits': ['Todos los registros que cumplen el filtro; no es una muestra aleatoria del universo.',
                   'La selección incluye solo detección por tránsito hasta 2024 con radio y periodo positivos sin banderas de límite.',
                   'Una fila representa un planeta. Las anfitrionas se repiten; sus parámetros pueden diferir entre referencias de planetas.',
                   'Los valores publicados pueden revisarse después de la fecha de corte del descubrimiento.'],
    },
    'poblacion': {
        'url': 'https://api.worldbank.org/v2/country/COL;PER;ECU/indicator/SP.POP.TOTL?date=2000:2024&format=json&per_page=1000',
        'documentation': 'https://data.worldbank.org/indicator/SP.POP.TOTL?locations=CO-PE-EC',
        'fields': ['pais', 'codigo', 'anio', 'poblacion'], 'numeric': ['anio', 'poblacion'],
        'nullable': ['poblacion'], 'key': ['codigo', 'anio'],
        'units': {'pais': 'nombre', 'codigo': 'ISO 3166-1 alfa-3', 'anio': 'año calendario', 'poblacion': 'personas'},
        'license': 'Creative Commons Attribution 4.0 (CC BY 4.0), según la ficha del indicador.',
        'license_url': 'https://data.worldbank.org/indicator/SP.POP.TOTL?locations=CO',
        'filters': 'SP.POP.TOTL; Colombia, Perú y Ecuador; años 2000–2024; 75 combinaciones país-año esperadas.',
        'transformation': 'Se extraen countryiso3code, date y value; se traducen únicamente los nombres de los tres países. Se ordena por código y año; los nulos se conservan vacíos.',
        'limits': ['Las estimaciones y series históricas pueden revisarse.',
                   'No sumar valores de distintos años ni tratar las cifras como microdatos individuales.',
                   'Las diferencias con DANE pueden corresponder a metodología, versión o fecha de referencia.'],
    },
    'clima': {
        'url': 'https://power.larc.nasa.gov/api/temporal/daily/point?' + urlencode({
            'parameters': 'T2M,PRECTOTCORR', 'community': 'AG', 'longitude': -74.0721,
            'latitude': 4.7110, 'start': '20240101', 'end': '20241231', 'format': 'JSON', 'time-standard': 'LST'}),
        'documentation': 'https://power.larc.nasa.gov/docs/services/api/temporal/daily/',
        'fields': ['fecha', 'temperatura_c', 'precipitacion_mm'],
        'numeric': ['temperatura_c', 'precipitacion_mm'], 'nullable': ['temperatura_c', 'precipitacion_mm'], 'key': ['fecha'],
        'units': {'fecha': 'día en hora solar local (LST)', 'temperatura_c': '°C', 'precipitacion_mm': 'mm acumulados en el día'},
        'license': 'Acceso gratuito a los datos NASA POWER; conservar atribución al proyecto y a sus fuentes. Revisar condiciones de la fuente para otros usos.',
        'license_url': 'https://power.larc.nasa.gov/',
        'filters': 'Punto de consulta 4.7110, -74.0721; Bogotá; 2024 completo; comunidad AG; días LST.',
        'transformation': 'T2M pasa a temperatura_c. PRECTOTCORR expresa mm/día: al integrar cada intervalo diario completo, el valor numérico equivale a mm del día. -999 se convierte en vacío. Se conserva LST.',
        'limits': ['Es una celda de rejilla de fuentes como MERRA2, no una estación urbana de Bogotá.',
                   'La elevación y la resolución espacial de la celda pueden diferir del sitio de interés.',
                   'Un solo año permite explorar estacionalidad; no demuestra una tendencia climática de largo plazo.',
                   'LST es hora solar local; no equivale al huso horario civil de Colombia.'],
    },
    'sismos': {
        'url': 'https://earthquake.usgs.gov/fdsnws/event/1/query?' + urlencode({
            'format': 'geojson', 'starttime': '2024-01-01T00:00:00', 'endtime': '2024-01-31T23:59:59.999',
            'minmagnitude': 5, 'eventtype': 'earthquake', 'orderby': 'time-asc', 'limit': 20000}),
        'documentation': 'https://earthquake.usgs.gov/fdsnws/event/1/',
        'fields': ['id', 'fecha', 'latitud', 'longitud', 'profundidad_km', 'magnitud', 'tipo_magnitud'],
        'numeric': ['latitud', 'longitud', 'profundidad_km', 'magnitud'], 'nullable': [], 'key': ['id'],
        'units': {'id': 'identificador USGS', 'fecha': 'instante UTC', 'latitud': 'grados', 'longitud': 'grados',
                  'profundidad_km': 'km', 'magnitud': 'escala indicada por tipo_magnitud', 'tipo_magnitud': 'categoría'},
        'license': 'Datos del catálogo USGS de acceso público; citar USGS y conservar la procedencia y fecha. Revisar créditos específicos de productos asociados.',
        'license_url': 'https://www.usgs.gov/information-policies-and-instructions/copyrights-and-credits',
        'filters': 'Catálogo global; enero de 2024 en UTC; tipo earthquake; magnitud ≥5; máximo técnico 20.000.',
        'transformation': 'Se transforma GeoJSON a tabla: coordinates[0] es longitud, [1] latitud y [2] profundidad. Tiempo Unix en milisegundos se convierte a ISO 8601 UTC.',
        'limits': ['Se rechaza una respuesta que alcance 20.000 filas para evitar truncamiento silencioso.',
                   'La magnitud tiene escala logarítmica; no mide directamente intensidad ni daños.',
                   'Los tipos de magnitud y las revisiones pueden variar. El catálogo no predice sismos futuros.'],
    },
}


def finite_number(value, nullable=False, missing=None):
    """Convierte números sin confundir faltantes, infinitos, booleanos y ceros."""
    if value is None or value == '' or (missing is not None and value == missing):
        if nullable:
            return ''
        raise ValueError('Falta un valor numérico obligatorio.')
    if isinstance(value, bool):
        raise ValueError('Un booleano no es una observación numérica.')
    number = float(value)
    if not math.isfinite(number):
        raise ValueError('La observación debe ser finita.')
    return int(number) if number.is_integer() else number


def normalize(name, payload):
    """Normaliza respuestas oficiales sin inventar observaciones."""
    if name not in DATASETS:
        raise ValueError('Conjunto desconocido.')
    config = DATASETS[name]
    rows = []
    if name == 'exoplanetas':
        for source in payload:
            row = {field: source[field] for field in config['fields']}
            for field in config['numeric']:
                row[field] = finite_number(row[field], field in config['nullable'])
            if row['pl_rade'] <= 0 or row['pl_orbper'] <= 0 or row['discoverymethod'] != 'Transit':
                raise ValueError('Un planeta no cumple el filtro de la consulta.')
            rows.append(row)
    elif name == 'poblacion':
        if payload[0]['pages'] != 1 or len(payload[1]) != int(payload[0]['total']):
            raise ValueError('Respuesta paginada incompleta; recorre todas las páginas antes de ampliar el conjunto.')
        names = {'COL': 'Colombia', 'PER': 'Perú', 'ECU': 'Ecuador'}
        for source in payload[1]:
            rows.append({'pais': names[source['countryiso3code']], 'codigo': source['countryiso3code'],
                         'anio': int(source['date']), 'poblacion': finite_number(source['value'], True)})
    elif name == 'clima':
        if payload['header']['time_standard'] != 'LST':
            raise ValueError('La consulta requiere días LST.')
        parameters = payload['properties']['parameter']
        if set(parameters['T2M']) != set(parameters['PRECTOTCORR']):
            raise ValueError('Las dos variables no cubren las mismas fechas.')
        for day in parameters['T2M']:
            rows.append({'fecha': datetime.strptime(day, '%Y%m%d').date().isoformat(),
                         'temperatura_c': finite_number(parameters['T2M'][day], True, payload['header']['fill_value']),
                         'precipitacion_mm': finite_number(parameters['PRECTOTCORR'][day], True, payload['header']['fill_value'])})
    else:
        if len(payload['features']) >= 20000:
            raise ValueError('Posible truncamiento del catálogo: divide la consulta por periodos.')
        for feature in payload['features']:
            properties = feature['properties']
            coordinates = feature['geometry']['coordinates']
            stamp = datetime.fromtimestamp(finite_number(properties['time']) / 1000, timezone.utc)
            rows.append({'id': feature['id'], 'fecha': stamp.isoformat(timespec='milliseconds').replace('+00:00', 'Z').replace('.000Z', 'Z'),
                         'latitud': finite_number(coordinates[1]), 'longitud': finite_number(coordinates[0]),
                         'profundidad_km': finite_number(coordinates[2]), 'magnitud': finite_number(properties['mag']),
                         'tipo_magnitud': properties['magType']})
    if not rows:
        raise ValueError('La fuente no devolvió observaciones.')
    keys = [tuple(row[field] for field in config['key']) for row in rows]
    if len(set(keys)) != len(rows):
        raise ValueError('Claves duplicadas; no se eliminan registros sin revisión.')
    rows.sort(key=lambda row: tuple(row[field] for field in config['key']))
    return rows


def write_csv(path, fields, rows):
    """Escribe CSV UTF-8 con encabezado único y sin columnas de índice."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', encoding='utf-8', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, lineterminator='\n')
        writer.writeheader()
        writer.writerows(rows)


def read_csv(path, name):
    """Lee el contrato publicado y recupera tipos numéricos explícitos."""
    config = DATASETS[name]
    with Path(path).open(encoding='utf-8-sig', newline='') as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames != config['fields']:
            raise ValueError(f'Columnas inesperadas en {name}.')
        rows = list(reader)
    for row in rows:
        for field in config['numeric']:
            row[field] = finite_number(row[field], field in config['nullable'])
    return rows


def fetch_bytes(url, opener=urlopen, max_bytes=10_000_000):
    """Descarga HTTPS acotada. El parámetro opener permite probar el límite de E/S."""
    if not url.startswith('https://'):
        raise ValueError('La fuente debe usar HTTPS.')
    request = Request(url, headers={'User-Agent': 'Guia-IA-Administrativos/analisis-educativo'})
    with opener(request, timeout=90) as response:
        raw = response.read(max_bytes + 1)
    if len(raw) > max_bytes:
        raise ValueError('La descarga supera 10 MB o el límite configurado; reduce la consulta.')
    return raw


def prepare_dataset(name, raw, base, timestamp):
    """Conserva el original, transforma y registra procedencia e integridad."""
    base = Path(base)
    config = DATASETS[name]
    payload = json.loads(raw)
    rows = normalize(name, payload)
    original = base / 'originales' / f'{name}.json'
    original.parent.mkdir(parents=True, exist_ok=True)
    original.write_bytes(raw)
    target = base / 'datos' / f'{name}.csv'
    write_csv(target, config['fields'], rows)
    metadata = {}
    if name == 'clima':
        metadata = {'time_standard': payload['header']['time_standard'],
                    'coordenadas_celda': payload['geometry']['coordinates'],
                    'fuentes': payload['header'].get('sources', []),
                    'unidades_originales': payload.get('parameters', {})}
    elif name == 'poblacion':
        metadata = {'actualizacion_fuente': payload[0].get('lastupdated'), 'paginas': payload[0]['pages']}
    return {'id': name, 'archivo': f'datos/{name}.csv', 'original': f'originales/{name}.json',
            'url_consulta': config['url'], 'url_documentacion': config['documentation'],
            'fecha_descarga': timestamp, 'filas': len(rows), 'columnas': config['fields'],
            'unidades': config['units'], 'sha256': hashlib.sha256(target.read_bytes()).hexdigest(),
            'sha256_original': hashlib.sha256(raw).hexdigest(), 'licencia': config['license'],
            'url_licencia': config['license_url'], 'filtros': config['filters'],
            'transformacion': config['transformation'], 'limites': config['limits'], 'metadatos': metadata}


def refresh(base=BASE, fetcher=fetch_bytes):
    """Actualiza las cuatro fuentes solo por una acción explícita de red."""
    base = Path(base)
    timestamp = datetime.now(timezone.utc).isoformat(timespec='seconds')
    # Descargar y validar todo antes de reemplazar la versión publicada.
    fetched = {name: fetcher(config['url']) for name, config in DATASETS.items()}
    for name, raw in fetched.items():
        normalize(name, json.loads(raw))
    manifest = {'version': 1, 'descripcion': 'Datos públicos reales de fuentes oficiales para aprendizaje; no son registros sintéticos.',
                'datasets': {name: prepare_dataset(name, raw, base, timestamp) for name, raw in fetched.items()}}
    manifest_path = base / 'fuentes.json'
    if manifest_path.exists():
        previous = json.loads(manifest_path.read_text(encoding='utf-8'))
        if 'ampliaciones' in previous:
            manifest['ampliaciones'] = previous['ampliaciones']
    (base / 'fuentes.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    write_dictionary(base, manifest)
    return manifest


def write_dictionary(base, manifest):
    """Documenta las variables básicas y las ampliaciones presentes."""
    base = Path(base)
    dictionary = []
    for name, config in DATASETS.items():
        for field in config['fields']:
            dictionary.append({'conjunto': name, 'columna': field, 'unidad': config['units'][field],
                               'tipo': 'número' if field in config['numeric'] else 'texto',
                               'faltante': 'celda vacía' if field in config['nullable'] else 'no permitido',
                               'fuente': config['documentation']})
    for name, item in manifest.get('ampliaciones', {}).items():
        for field in item['columnas']:
            dictionary.append({'conjunto': name, 'columna': field, 'unidad': item['unidades'][field],
                               'tipo': 'número', 'faltante': 'no permitido', 'fuente': item['url_consulta']})
    write_csv(base / 'diccionario.csv', ['conjunto', 'columna', 'unidad', 'tipo', 'faltante', 'fuente'], dictionary)


def validate_bundle(base=BASE):
    """Verifica hashes, contratos, claves y equivalencia con originales offline."""
    base = Path(base)
    manifest = json.loads((base / 'fuentes.json').read_text(encoding='utf-8'))
    if set(manifest['datasets']) != set(DATASETS):
        raise ValueError('El manifiesto debe contener los cuatro conjuntos.')
    counts = {}
    for name, config in DATASETS.items():
        item = manifest['datasets'][name]
        path = base / item['archivo']
        original = base / item['original']
        if hashlib.sha256(path.read_bytes()).hexdigest() != item['sha256']:
            raise ValueError(f'SHA256 del CSV no coincide: {name}.')
        if hashlib.sha256(original.read_bytes()).hexdigest() != item['sha256_original']:
            raise ValueError(f'SHA256 del original no coincide: {name}.')
        rows = read_csv(path, name)
        expected = normalize(name, json.loads(original.read_bytes()))
        if rows != expected or len(rows) != item['filas'] or item['columnas'] != config['fields']:
            raise ValueError(f'La transformación no coincide con el original: {name}.')
        if item['url_consulta'] != config['url']:
            raise ValueError(f'La URL no coincide con la consulta reproducible: {name}.')
        counts[name] = len(rows)
    return {'estado': 'verificado', 'filas': counts}


def summarize(name, rows):
    """Calcula resultados descriptivos; no atribuye causalidad ni predice."""
    if name not in DATASETS or not rows:
        raise ValueError('Se requieren un conjunto conocido y observaciones.')
    result = {'filas': len(rows)}
    if name == 'exoplanetas':
        result.update(anfitrionas_unicas=len({row['hostname'] for row in rows}),
                      mediana_radio_terrestre=statistics.median(row['pl_rade'] for row in rows),
                      mediana_periodo_dias=statistics.median(row['pl_orbper'] for row in rows),
                      planetas_sin_temperatura_estelar=sum(row['st_teff'] == '' for row in rows))
    elif name == 'poblacion':
        countries = {}
        for code in sorted({row['codigo'] for row in rows}):
            series = sorted((row for row in rows if row['codigo'] == code and row['poblacion'] != ''), key=lambda row: row['anio'])
            if not series:
                countries[code] = {'estado': 'sin observaciones disponibles'}
                continue
            first, last = series[0], series[-1]
            countries[code] = {'pais': first['pais'], 'anio_inicial': first['anio'], 'anio_final': last['anio'],
                               'poblacion_inicial': first['poblacion'], 'poblacion_final': last['poblacion'],
                               'cambio_porcentual': (last['poblacion'] / first['poblacion'] - 1) * 100 if first['poblacion'] else None}
        result['paises'] = countries
    elif name == 'clima':
        temperatures = [row['temperatura_c'] for row in rows if row['temperatura_c'] != '']
        rainfall = [row['precipitacion_mm'] for row in rows if row['precipitacion_mm'] != '']
        result.update(temperatura_media_c=statistics.mean(temperatures) if temperatures else None,
                      precipitacion_total_disponible_mm=math.fsum(rainfall) if rainfall else None,
                      dias_sin_temperatura=len(rows) - len(temperatures), dias_sin_precipitacion=len(rows) - len(rainfall))
    else:
        result.update(magnitud_maxima=max(row['magnitud'] for row in rows),
                      mediana_profundidad_km=statistics.median(row['profundidad_km'] for row in rows),
                      tipos_magnitud=sorted({row['tipo_magnitud'] for row in rows}))
    return result


def analyze(base=BASE, output=None):
    """Escribe un resumen trazable sin actualizar ni alterar las fuentes."""
    base = Path(base)
    validate_bundle(base)
    output = Path(output) if output is not None else base / 'codigo/resultados'
    output.mkdir(parents=True, exist_ok=True)
    result = {name: summarize(name, read_csv(base / 'datos' / f'{name}.csv', name)) for name in DATASETS}
    (output / 'resumen.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    return result


def create_figures(base=BASE, output=None, only=None):
    """Crea figuras científicas PNG/SVG, con ejes, unidades y selección explícita."""
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    base = Path(base)
    validate_bundle(base)
    output = Path(output) if output is not None else base / 'codigo/resultados'
    output.mkdir(parents=True, exist_ok=True)
    paths = []
    plt.rcParams.update({'font.size': 10, 'axes.spines.top': False, 'axes.spines.right': False,
                         'svg.fonttype': 'none'})
    for name in ([only] if only else DATASETS):
        rows = read_csv(base / 'datos' / f'{name}.csv', name)
        fig, axes = plt.subplots(1, 2, figsize=(12, 4.8), layout='constrained')
        if name == 'exoplanetas':
            axes[0].scatter([r['pl_orbper'] for r in rows], [r['pl_rade'] for r in rows], s=10, alpha=.4, color='#156b8a')
            axes[0].set(xscale='log', yscale='log', xlabel='Periodo orbital (días)', ylabel='Radio (radios terrestres)', title='Planetas detectados por tránsito')
            axes[1].hist([r['pl_rade'] for r in rows], bins=25, color='#156b8a', edgecolor='white')
            axes[1].set(xlabel='Radio (radios terrestres)', ylabel='Número de planetas', title='Distribución en la selección publicada')
            fig.suptitle(f'NASA Exoplanet Archive · {len(rows):,} planetas · descubrimiento hasta 2024')
        elif name == 'poblacion':
            for code, color in [('COL', '#156b8a'), ('ECU', '#a45b13'), ('PER', '#75439a')]:
                series = sorted((r for r in rows if r['codigo'] == code and r['poblacion'] != ''), key=lambda r: r['anio'])
                if not series:
                    continue
                axes[0].plot([r['anio'] for r in series], [r['poblacion']/1e6 for r in series], label=series[0]['pais'], color=color)
                if series[0]['poblacion']:
                    axes[1].plot([r['anio'] for r in series], [r['poblacion']/series[0]['poblacion']*100 for r in series], label=series[0]['pais'], color=color)
            axes[0].set(xlabel='Año', ylabel='Población (millones)', title='Tamaño de la población')
            axes[1].set(xlabel='Año', ylabel='Índice (primer año disponible = 100)', title='Cambio relativo con la misma base')
            axes[0].legend(frameon=False)
            axes[1].legend(frameon=False)
            fig.suptitle('Banco Mundial · SP.POP.TOTL · estimaciones publicadas')
        elif name == 'clima':
            rows.sort(key=lambda r: r['fecha'])
            dates = [datetime.fromisoformat(r['fecha']) for r in rows]
            axes[0].plot(dates, [float('nan') if r['temperatura_c'] == '' else r['temperatura_c'] for r in rows], color='#a45b13', linewidth=.9)
            axes[1].bar(dates, [float('nan') if r['precipitacion_mm'] == '' else r['precipitacion_mm'] for r in rows], color='#156b8a', width=1)
            axes[0].set(xlabel='Fecha (LST)', ylabel='Temperatura a 2 m (°C)', title='Promedio diario de temperatura')
            axes[1].set(xlabel='Fecha (LST)', ylabel='Precipitación diaria (mm)', title='Precipitación de cada día')
            for ax in axes:
                ax.tick_params(axis='x', labelrotation=35)
            fig.suptitle('NASA POWER · celda consultada para Bogotá · 2024 · datos de rejilla')
        else:
            plotted = axes[0].scatter([r['longitud'] for r in rows], [r['latitud'] for r in rows],
                                     c=[r['profundidad_km'] for r in rows], s=18, cmap='cividis', alpha=.8)
            fig.colorbar(plotted, ax=axes[0], label='Profundidad (km)')
            axes[0].set(xlim=(-180, 180), ylim=(-90, 90), xlabel='Longitud (°)', ylabel='Latitud (°)', title='Ubicación de eventos (sin mapa base)')
            axes[1].hist([r['profundidad_km'] for r in rows], bins=20, color='#156b8a', edgecolor='white')
            axes[1].set(xlabel='Profundidad (km)', ylabel='Número de eventos', title='Distribución de profundidades')
            fig.suptitle('USGS · enero de 2024 · magnitud ≥5 · catálogo global')
        for ax in axes:
            ax.grid(alpha=.15)
        target = output / f'{name}.png'
        fig.savefig(target, dpi=160)
        fig.savefig(target.with_suffix('.svg'))
        plt.close(fig)
        paths.append(target)
    return paths


def normalize_dane(header, source):
    """Extrae edades simples y comprueba hombres + mujeres = total publicado."""
    ages = {}
    for column, label in enumerate(header):
        match = re.match(r'^(Hombres|Mujeres) (\d+) añ', str(label))
        if match:
            age = int(match.group(2))
            sex = match.group(1).lower()
            ages.setdefault(age, {'edad': age})[sex] = finite_number(source[column])
    rows = [ages[age] for age in sorted(ages)]
    if not rows or any(set(row) != {'edad', 'hombres', 'mujeres'} for row in rows):
        raise ValueError('No se encontraron edades con ambos sexos.')
    men = sum(row['hombres'] for row in rows)
    women = sum(row['mujeres'] for row in rows)
    if men != source[5] or women != source[6] or men + women != source[4]:
        raise ValueError('Las edades no concilian con los totales publicados.')
    return rows


def refresh_dane(base=BASE, fetcher=fetch_bytes, workbook_loader=None):
    """Lee XLSX en memoria y conserva solo un extracto verificable de Santander."""
    if workbook_loader is None:
        from openpyxl import load_workbook
        workbook_loader = load_workbook
    base = Path(base)
    raw = fetcher(DANE_URL)
    workbook = workbook_loader(io.BytesIO(raw), read_only=True, data_only=True)
    sheet = workbook['PobDepartamentalxÁreaSexoEdad']
    header = list(next(sheet.iter_rows(min_row=9, max_row=9, values_only=True)))
    selected = [list(row) for row in sheet.iter_rows(min_row=10, values_only=True)
                if str(row[0]) == '68' and row[1] == 'Santander' and row[2] == 2024 and row[3] == 'Total']
    workbook.close()
    if len(selected) != 1:
        raise ValueError('Se requiere exactamente una fila de Santander, 2024, área Total.')
    source = selected[0]
    rows = normalize_dane(header, source)
    target = base / 'datos/piramide-santander-2024.csv'
    write_csv(target, ['edad', 'hombres', 'mujeres'], rows)
    original = base / 'originales/dane-santander-2024-extracto.json'
    original.parent.mkdir(parents=True, exist_ok=True)
    original.write_text(json.dumps({'hoja': 'PobDepartamentalxÁreaSexoEdad', 'fila_encabezado': 9,
                                   'encabezados': header, 'fila': source}, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    item = {'id': 'dane_santander', 'archivo': 'datos/piramide-santander-2024.csv',
            'original': 'originales/dane-santander-2024-extracto.json', 'url_consulta': DANE_URL,
            'fecha_descarga': datetime.now(timezone.utc).isoformat(timespec='seconds'),
            'sha256': hashlib.sha256(target.read_bytes()).hexdigest(),
            'sha256_original': hashlib.sha256(original.read_bytes()).hexdigest(),
            'sha256_xlsx_consultado': hashlib.sha256(raw).hexdigest(), 'bytes_xlsx_consultado': len(raw),
            'filas': len(rows), 'columnas': ['edad', 'hombres', 'mujeres'],
            'unidades': {'edad': 'años cumplidos; 100 representa 100 años y más', 'hombres': 'personas', 'mujeres': 'personas'},
            'total': source[4], 'hombres': source[5], 'mujeres': source[6],
            'filtros': 'Departamento 68, Santander; año 2024; área Total; encabezado fila 9.',
            'transformacion': 'Se extraen columnas Hombres/Mujeres por edad. Se comprueba la suma por sexo y ambos sexos contra la misma fila. No se suman áreas con Total.',
            'licencia': 'Fuente pública DANE. Se distribuye únicamente un extracto numérico con atribución; el XLSX completo se consulta en el enlace oficial y no se incluye en el paquete.',
            'url_documentacion': 'https://www.dane.gov.co/index.php/estadisticas-por-tema/demografia-y-poblacion/proyecciones-de-poblacion',
            'limites': ['Son proyecciones con base en CNPV 2018, no un conteo censal de 2024.',
                       'El grupo final representa 100 años y más. La figura agrupa edades en intervalos de cinco años.',
                       'La página publica el anexo el 8 de agosto de 2025; la hoja PPED declara actualización del 30 de julio de 2025.',
                       'Se conserva un extracto de las celdas consultadas; el archivo XLSX completo no se redistribuye.']}
    manifest_path = base / 'fuentes.json'
    manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
    manifest.setdefault('ampliaciones', {})['dane_santander'] = item
    write_dictionary(base, manifest)
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    return item


def validate_dane(base=BASE):
    """Comprueba el extracto DANE y la tabla derivada sin conexión a internet."""
    base = Path(base)
    item = json.loads((base / 'fuentes.json').read_text(encoding='utf-8'))['ampliaciones']['dane_santander']
    original = base / item['original']
    target = base / item['archivo']
    if hashlib.sha256(original.read_bytes()).hexdigest() != item['sha256_original'] or hashlib.sha256(target.read_bytes()).hexdigest() != item['sha256']:
        raise ValueError('SHA256 de DANE no coincide.')
    extract = json.loads(original.read_text(encoding='utf-8'))
    rows = normalize_dane(extract['encabezados'], extract['fila'])
    with target.open(encoding='utf-8', newline='') as handle:
        saved = [{key: int(value) for key, value in row.items()} for row in csv.DictReader(handle)]
    if rows != saved or len(rows) != item['filas']:
        raise ValueError('La pirámide no coincide con las celdas originales.')
    return {'total': sum(row['hombres'] + row['mujeres'] for row in rows),
            'hombres': sum(row['hombres'] for row in rows), 'mujeres': sum(row['mujeres'] for row in rows)}


def create_pyramid(base=BASE, output=None):
    """Grafica grupos de edad; el signo negativo solo determina la posición."""
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.ticker import FuncFormatter
    base = Path(base)
    validate_dane(base)
    output = Path(output) if output is not None else base / 'codigo/resultados'
    output.mkdir(parents=True, exist_ok=True)
    groups = {}
    with (base / 'datos/piramide-santander-2024.csv').open(encoding='utf-8', newline='') as handle:
        for row in csv.DictReader(handle):
            age = int(row['edad'])
            group = age // 5 * 5
            groups.setdefault(group, [0, 0])
            groups[group][0] += int(row['hombres'])
            groups[group][1] += int(row['mujeres'])
    ages = sorted(groups)
    fig, ax = plt.subplots(figsize=(8, 8), layout='constrained')
    positions = list(range(len(ages)))
    ax.barh(positions, [-groups[age][0] for age in ages], color='#156b8a', label='Hombres')
    ax.barh(positions, [groups[age][1] for age in ages], color='#a45b13', label='Mujeres')
    ax.set_yticks(positions, ['100 y más' if age == 100 else f'{age}–{age+4}' for age in ages])
    ax.xaxis.set_major_formatter(FuncFormatter(lambda value, position: f'{abs(value):,.0f}'))
    ax.set(xlabel='Personas (valores absolutos)', ylabel='Grupo de edad (años)',
           title='Santander · proyección DANE para 2024\nÁrea total · base CNPV 2018 · actualización 2025')
    ax.axvline(0, color='#555555', linewidth=.8)
    ax.legend(frameon=False)
    target = output / 'piramide-santander-2024.png'
    fig.savefig(target, dpi=150)
    fig.savefig(target.with_suffix('.svg'))
    plt.close(fig)
    return target


def main(argv=None):
    """Interfaz con acceso a la red solo en refrescar y refrescar-dane."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('accion', choices=['validar', 'analizar', 'figuras', 'refrescar', 'refrescar-dane', 'validar-dane', 'piramide'])
    parser.add_argument('--base', type=Path, default=BASE)
    parser.add_argument('--salida', type=Path)
    args = parser.parse_args(argv)
    if args.accion == 'refrescar-dane':
        result = refresh_dane(args.base, fetcher=fetch_bytes)
    elif args.accion == 'validar-dane':
        result = validate_dane(args.base)
    elif args.accion == 'piramide':
        result = str(create_pyramid(args.base, args.salida))
    elif args.accion == 'refrescar':
        result = refresh(args.base, fetcher=fetch_bytes)
    elif args.accion == 'validar':
        result = validate_bundle(args.base)
    elif args.accion == 'analizar':
        result = analyze(args.base, args.salida)
    else:
        result = [str(path) for path in create_figures(args.base, args.salida)]
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
